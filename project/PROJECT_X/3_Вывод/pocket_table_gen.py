"""Макрос 1. Формирование Excel-таблицы состава пакета."""

from __future__ import annotations

from pathlib import Path
import subprocess
import sys
import tkinter as tk
from tkinter import filedialog, messagebox

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except ModuleNotFoundError as error:
    if error.name != "openpyxl":
        raise
    Workbook = None
    Alignment = Border = Font = PatternFill = Side = None


OPENPYXL_REQUIREMENT = "openpyxl>=3.1,<4"
PYWIN32_REQUIREMENT = "pywin32>=306"


def ensure_dependencies() -> str | None:
    """
    Проверяет наличие openpyxl и при необходимости устанавливает зависимости.

    Установка выполняется через тот же интерпретатор Python, которым запущен
    текущий макрос. Поэтому Visual Studio Python и системный Python не смешиваются.
    Возвращает текст ошибки только если автоматическая установка не удалась.
    """
    global Workbook, Alignment, Border, Font, PatternFill, Side

    if Workbook is not None:
        return None

    install_command = [
        sys.executable,
        "-m",
        "pip",
        "install",
        OPENPYXL_REQUIREMENT,
    ]

    try:
        subprocess.check_call(install_command)
    except (OSError, subprocess.CalledProcessError) as error:
        command_text = " ".join(f'"{part}"' if " " in part else part for part in install_command)
        return (
            "Не удалось автоматически установить зависимости.\n\n"
            f"Команда: {command_text}\n\n"
            f"Подробности: {error}"
        )

    try:
        from openpyxl import Workbook as workbook_class
        from openpyxl.styles import (
            Alignment as alignment_class,
            Border as border_class,
            Font as font_class,
            PatternFill as pattern_fill_class,
            Side as side_class,
        )
    except ModuleNotFoundError as error:
        return f"Зависимость установлена, но openpyxl не загрузился: {error}"

    Workbook = workbook_class
    Alignment = alignment_class
    Border = border_class
    Font = font_class
    PatternFill = pattern_fill_class
    Side = side_class
    return None
# ============================================================================
# НАСТРОЙКИ МАКРОСА
#
# Все параметры, которые обычно требуется менять при адаптации макроса,
# находятся в этом блоке:
#
# PROJECT_ROOT
#     Корень проекта. По умолчанию это папка, из которой запущен скрипт.
#
# IGNORED_FOLDER_NAMES
#     Папки, которые не попадут в таблицу и не будут просматриваться.
#
# IGNORED_FILE_EXTENSIONS
#     Расширения файлов, которые нужно исключить из списка.
#
# DEFAULT_STATUS
#     Начальное логическое значение во втором столбце: False — не включать.
#
# FONT_NAME, FONT_SIZE
#     Шрифт и размер шрифта итоговой таблицы.
# ============================================================================
PROJECT_ROOT = Path.cwd().resolve()


IGNORED_FOLDER_NAMES = {"0_Архив", "1_Пакеты на отправку"}


IGNORED_FILE_EXTENSIONS = {".py", ".xlsx", ".xlsm"}


DEFAULT_STATUS = True


FONT_NAME = "ISOCPEUR"
FONT_SIZE = 14


# ============================================================================
# СБОР ДАННЫХ ИЗ СТРУКТУРЫ ПРОЕКТА
# ============================================================================
def collect_manifest(root: Path) -> list[tuple[str, list[str]]]:
    """
    Формирует данные для будущей таблицы.

    В таблицу попадают:
    - только папки первого уровня относительно root;
    - файлы на любой глубине внутри этих папок;
    - только файлы, расширения которых не указаны в списке исключений.

    Названия папок и файлов сортируются по имени, без учета регистра.
    """
    result = []

    for folder in sorted(root.iterdir(), key=lambda item: item.name.casefold()):

        if not folder.is_dir() or folder.name in IGNORED_FOLDER_NAMES:
            continue

        files = sorted(
            (
                path.name
                for path in folder.rglob("*")
                if path.is_file()
                and path.suffix.casefold() not in IGNORED_FILE_EXTENSIONS
            ),
            key=str.casefold,
        )

        result.append((folder.name, files))

    return result


# ============================================================================
# СОЗДАНИЕ И ФОРМАТИРОВАНИЕ EXCEL-КНИГИ
# ============================================================================
def build_workbook(package_name: str, manifest: list[tuple[str, list[str]]]) -> Workbook:
    """Создает книгу, заполняет состав и применяет оформление."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Состав пакета"
    sheet["A1"] = package_name
    sheet.merge_cells("A1:B1")

    # ------------------------------------------------------------------------
    # Заполняем строки:
    # строка без статуса во втором столбце обозначает главную папку;
    # строки со статусом обозначают файлы внутри последней папки.
    # ------------------------------------------------------------------------
    row = 2

    for folder_name, files in manifest:
        sheet.cell(row, 1, folder_name)
        sheet.cell(row, 2, None)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        row += 1

        for file_name in files:
            sheet.cell(row, 1, file_name)
            sheet.cell(row, 2, DEFAULT_STATUS)
            row += 1

    last_row = max(row - 1, 1)

    # ------------------------------------------------------------------------
    # Настраиваем цвета и границы.
    # ------------------------------------------------------------------------
    thin_side = Side(style="thin", color="000000")
    border = Border(
        left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
    )
    yellow_fill = PatternFill("solid", fgColor="FFF2CC")

    # ------------------------------------------------------------------------
    # Применяем общий стиль ко всем ячейкам активной области.
    # ------------------------------------------------------------------------
    for current_row in sheet.iter_rows(min_row=1, max_row=last_row, min_col=1, max_col=2):

        for cell in current_row:
            cell.font = Font(name=FONT_NAME, size=FONT_SIZE, bold=False)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
            cell.number_format = "@"

    sheet["A1"].font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True)

    # Строки главных папок выделяются жирным шрифтом и желтой заливкой.
    for current_row in range(2, last_row + 1):

        if sheet.cell(current_row, 2).value is None:
            sheet.cell(current_row, 1).font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True)
            sheet.cell(current_row, 1).fill = yellow_fill

    # ------------------------------------------------------------------------
    # Настраиваем размеры, выбор статуса и область печати.
    # ------------------------------------------------------------------------
    sheet.column_dimensions["A"].width = 100
    sheet.column_dimensions["B"].width = 25

    for current_row in range(1, last_row + 1):
        sheet.row_dimensions[current_row].height = 25

    # Ячейки B будут связаны с настоящими флажками Excel после сохранения.

    sheet.print_area = f"A1:B{last_row}"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_view.view = "pageBreakPreview"
    sheet.freeze_panes = "A2"

    return workbook


def get_file_rows(manifest: list[tuple[str, list[str]]]) -> list[int]:
    """Возвращает номера строк Excel, соответствующих файлам."""
    rows = []
    row_number = 2

    for _, files in manifest:
        row_number += 1

        for _ in files:
            rows.append(row_number)
            row_number += 1

    return rows


def ensure_checkbox_dependency() -> str | None:
    """Устанавливает pywin32, необходимый для управления Excel через COM."""
    try:
        import win32com.client
        return None
    except ModuleNotFoundError:
        pass

    install_command = [
        sys.executable,
        "-m",
        "pip",
        "install",
        PYWIN32_REQUIREMENT,
    ]

    try:
        subprocess.check_call(install_command)
    except (OSError, subprocess.CalledProcessError) as error:
        return f"Не удалось установить pywin32 для создания чекбоксов: {error}"

    return None


def add_excel_checkboxes(workbook_path: Path, file_rows: list[int]) -> None:
    """
    Добавляет классические интерактивные CheckBox-контролы через Excel COM.

    openpyxl не умеет создавать Form Controls, поэтому Excel открывается
    невидимо, флажки размещаются поверх ячеек B, а их состояния связываются
    с этими ячейками как TRUE/FALSE.
    """
    import win32com.client
    import pywintypes

    excel = win32com.client.DispatchEx("Excel.Application")
    workbook = None

    try:
        excel.Visible = False
        excel.DisplayAlerts = False
        workbook = excel.Workbooks.Open(str(workbook_path.resolve()))
        sheet = workbook.Worksheets("Состав пакета")

        if sheet.CheckBoxes().Count > 0:
            sheet.CheckBoxes().Delete()

        for row_number in file_rows:
            cell = sheet.Cells(row_number, 2)
            checkbox = sheet.CheckBoxes().Add(
                cell.Left + 2,
                cell.Top + 2,
                cell.Width - 4,
                cell.Height - 4,
            )
            checkbox.Caption = ""
            checkbox.LinkedCell = cell.Address
            # xlOn = 1: новые файлы по умолчанию включены в пакет.
            checkbox.Value = 1
            checkbox.Placement = 1

        workbook.Save()
    except (OSError, pywintypes.com_error) as error:
        raise RuntimeError(
            "Не удалось создать чекбоксы через Microsoft Excel. "
            "Проверьте, что Excel установлен и файл не открыт.\n"
            f"Подробности: {error}"
        ) from error
    finally:
        if workbook is not None:
            workbook.Close(SaveChanges=False)
        excel.Quit()


# ============================================================================
# ДИАЛОГ СОХРАНЕНИЯ
# ============================================================================
def choose_output_path() -> Path | None:
    """Показывает окно выбора места для сохранения таблицы."""
    root = tk.Tk()
    root.withdraw()

    selected = filedialog.asksaveasfilename(
        title="Сохранить таблицу состава пакета",
        initialdir=PROJECT_ROOT,
        defaultextension=".xlsx",
        filetypes=[("Excel workbook", "*.xlsx")],
    )

    root.destroy()

    return Path(selected) if selected else None


def show_result_window(title: str, text: str, success: bool) -> None:
    """Показывает короткое итоговое окно с цветом результата."""
    root = tk.Tk()
    root.withdraw()

    background_color = "#C6EFCE" if success else "#FFC7CE"
    text_color = "#006100" if success else "#9C0006"

    window = tk.Toplevel(root)
    window.title(title)
    window.configure(bg=background_color)
    window.resizable(False, False)

    label = tk.Label(
        window,
        text=text,
        bg=background_color,
        fg=text_color,
        padx=28,
        pady=22,
        font=("Segoe UI", 12, "bold"),
    )
    label.pack()

    button = tk.Button(window, text="OK", command=root.destroy, width=12)
    button.pack(pady=(0, 16))

    window.protocol("WM_DELETE_WINDOW", root.destroy)
    window.grab_set()
    root.mainloop()


# ============================================================================
# ТОЧКА ВХОДА
#
# Запуск без аргументов: открывается окно сохранения (удобно для double-click).
# Запуск с аргументом: путь к таблице передается через командную строку.
# ============================================================================
def main() -> int:
    dependency_error = ensure_dependencies()
    if dependency_error is not None:
        if len(sys.argv) > 1:
            print(dependency_error, file=sys.stderr)
        else:
            messagebox.showerror("Ошибка установки зависимостей", dependency_error)
        return 1

    output_path = Path(sys.argv[1]) if len(sys.argv) > 1 else choose_output_path()

    if output_path is None:
        return 0

    if output_path.suffix.casefold() != ".xlsx":
        output_path = output_path.with_suffix(".xlsx")

    try:
        # Основной пайплайн макроса:
        # 1. Сканируем папку запуска.
        # 2. Формируем структуру "папка -> файлы".
        # 3. Создаем и сохраняем Excel-таблицу.
        manifest = collect_manifest(PROJECT_ROOT)
        package_name = output_path.stem
        build_workbook(package_name, manifest).save(output_path)

        dependency_error = ensure_checkbox_dependency()
        if dependency_error is not None:
            raise RuntimeError(dependency_error)

        add_excel_checkboxes(output_path, get_file_rows(manifest))
    except (OSError, ValueError, RuntimeError) as error:
        if len(sys.argv) > 1:
            print(f"ОШИБКА ГЕНЕРАЦИИ\n{error}", file=sys.stderr)
        else:
            show_result_window("Ошибка генерации", "ОШИБКА ГЕНЕРАЦИИ", success=False)
        return 1

    if len(sys.argv) > 1:
        print(f"ТАБЛИЦА СГЕНЕРИРОВАНА\n{output_path}")
    else:
        show_result_window("Готово", "ТАБЛИЦА СГЕНЕРИРОВАНА", success=True)

    return 0


if __name__ == "__main__":
    sys.exit(main())